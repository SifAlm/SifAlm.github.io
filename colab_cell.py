import importlib
import json
import os
import random
import re
import subprocess
import sys
from collections import defaultdict
from dataclasses import dataclass, field, asdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from urllib.parse import urljoin, urlparse


###############################################################################
# Dependency bootstrap                                                         #
###############################################################################

def ensure_dependencies() -> None:
    """Install third-party packages required for the scraper."""

    required = {
        "playwright": "playwright",
        "pandas": "pandas",
        "openpyxl": "openpyxl",
        "bs4": "beautifulsoup4",
        "lxml": "lxml",
        "tqdm": "tqdm",
        "fake_useragent": "fake_useragent",
        "nest_asyncio": "nest_asyncio",
        "requests": "requests",
        "PIL": "pillow",
    }
    missing: List[str] = []
    for module_name, package_name in required.items():
        try:
            importlib.import_module(module_name)
        except ImportError:
            missing.append(package_name)
    if missing:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", *missing])
    subprocess.check_call([sys.executable, "-m", "playwright", "install", "--with-deps", "chromium"])


ensure_dependencies()

import asyncio
import contextlib
import io
import time

import nest_asyncio
import pandas as pd
import requests
from bs4 import BeautifulSoup
from fake_useragent import UserAgent
from PIL import Image, ImageDraw
from playwright.async_api import (
    Browser,
    BrowserContext,
    ElementHandle,
    Locator,
    Page,
    TimeoutError as PlaywrightTimeout,
    async_playwright,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

nest_asyncio.apply()


###############################################################################
# Global configuration                                                         #
###############################################################################

BASE_URL_DEFAULT = "https://www.justlife.com/en-AE"
BASE_URL = os.environ.get("BASE_URL", BASE_URL_DEFAULT)
OUTPUT_DIR = Path("justlife_intel")
OUTPUT_DIR.mkdir(exist_ok=True)
SHOTS_DIR = OUTPUT_DIR / "shots"
SHOTS_DIR.mkdir(exist_ok=True)

SERVICE_URL_TARGET = 90
CITY_PRIORITY = ["dubai", "abu-dhabi", "sharjah"]
MAX_COMBOS_PER_SERVICE = int(os.environ.get("MAX_COMBOS", os.environ.get("MAX_COMBOS_PER_SERVICE", "60")))
DAYS_AHEAD = int(os.environ.get("DAYS_AHEAD", os.environ.get("DAYS", "14")))
HEADLESS = os.environ.get("HEADLESS", "1") not in {"0", "false", "False"}
USER_AGENT = UserAgent().chrome

CTA_TEXT_PATTERNS = [
    "book now",
    "book",
    "schedule",
    "get started",
    "continue",
    "checkout",
    "select date",
    "select time",
    "start",
]
BOOKING_WIDGET_PROBES = [
    "[data-testid*='hour']",
    "[data-testid*='duration']",
    "[data-testid*='professional']",
    "[data-testid*='bed']",
    "[data-testid*='bath']",
    "[data-testid*='package']",
    "[data-testid*='quantity']",
    "[data-testid*='area']",
    "[data-testid*='slot']",
    "select",
    "button[aria-pressed]",
    "input[type='number']",
    "input[type='radio']",
]
PRICE_SELECTORS = [
    "[data-testid*='grand']",
    "[data-testid*='total']",
    "[data-testid*='price']",
    "[class*='grand']",
    "[class*='total']",
    "[class*='price']",
]
SUMMARY_HINTS = [
    "summary",
    "total",
    "vat",
    "fee",
    "breakdown",
]
CONTROL_LABEL_HINTS = {
    "HOURS_PROS": ["hour", "duration", "cleaner", "professional", "maid"],
    "ROOM_MATERIAL": ["bed", "bath", "room", "apartment", "villa", "material"],
    "UNIT_QUANTITY": ["unit", "piece", "item", "sofa", "mattress", "curtain", "ac"],
    "AREA_BASED": ["area", "sqft", "square", "size"],
    "FLAT_PACKAGE": ["package", "plan", "combo", "bundle", "standard", "premium"],
}
SERVICE_KEYWORDS = re.compile(
    r"(clean|maid|deep|move|villa|apartment|sofa|carpet|mattress|curtain|ac|pest|laundry|salon|spa|wax|brow|lash|nail|men|women|pet|plumb|electric|handyman|doctor|nurse|lab|physio|iv|oxygen|pack)",
    re.IGNORECASE,
)
EXCLUDE_PATTERNS = re.compile(
    r"/(?:dubai|abu-dhabi|sharjah|ajman)(?:/|$)|/my-account/|/faq|/terms|/privacy|/career|/sitemap|/blog|/what-|/how-",
    re.IGNORECASE,
)

RANDOM_SLEEP = (0.45, 1.05)


###############################################################################
# Data structures                                                              #
###############################################################################


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


@dataclass
class ServiceCandidate:
    url: str
    source: str
    is_checkout: bool = False
    city: Optional[str] = None


@dataclass
class ServiceRecord:
    page_id: str
    parent_page_id: Optional[str]
    level: str
    url: str
    page_title_text: str
    detected_model_type: str
    has_addons: bool
    has_slots: bool
    notes: str


@dataclass
class OptionRecord:
    page_id: str
    control_group: str
    option_code: str
    option_label: str
    option_type: str
    is_default: bool
    sort_index: int


@dataclass
class PriceRecord:
    page_id: str
    service_name: str
    model_type: str
    hours: Optional[str]
    pros: Optional[str]
    bedrooms: Optional[str]
    bathrooms: Optional[str]
    kitchen_package: Optional[str]
    appliance_combo: Optional[str]
    package_name: Optional[str]
    duration: Optional[str]
    offer_state: str
    offer_name: Optional[str]
    discount_type: Optional[str]
    discount_value: Optional[float]
    base_price: Optional[float]
    fees_total: Optional[float]
    vat_amount: Optional[float]
    vat_percent: Optional[float]
    grand_total: Optional[float]
    currency: Optional[str]
    pricing_source_url: str
    collected_at: str
    qa_flags: str
    screenshot_ref: Optional[str]


@dataclass
class AddonRecord:
    page_id: str
    addon_name: str
    addon_description: str
    addon_price: Optional[float]
    currency: Optional[str]
    is_recommended: bool
    is_required: bool
    seen_on_step: str
    pricing_source_url: str
    collected_at: str


@dataclass
class AvailabilityRecord:
    page_id: str
    date_str: str
    time_slot_label: str
    slot_status: str
    lead_time_days: int
    collected_at: str


@dataclass
class OfferRecord:
    page_id: str
    offer_name: str
    offer_badge_text: str
    description: str
    eligibility: str
    discount_type: Optional[str]
    discount_value: Optional[float]
    stacking_rules: str
    collected_at: str


@dataclass
class ElementRecord:
    page_id: str
    selector: str
    element_role: str
    raw_text: str
    normalized_text: str


@dataclass
class ErrorRecord:
    page_id: str
    url: str
    context: str
    message: str
    screenshot: Optional[str] = None


###############################################################################
# Utility helpers                                                              #
###############################################################################


def slugify(value: str) -> str:
    value = re.sub(r"[^a-zA-Z0-9-_]+", "-", value.strip().lower())
    value = re.sub(r"-+", "-", value)
    return value.strip("-") or "page"


def canonicalize_url(url: str) -> str:
    if not url:
        return ""
    url = url.split("#", 1)[0].strip()
    base_parts = urlparse(BASE_URL)
    joined = urljoin(BASE_URL, url)
    parsed = urlparse(joined)

    scheme = parsed.scheme or base_parts.scheme or "https"
    netloc = parsed.netloc or base_parts.netloc
    path = re.sub(r"/en-AE/(en-AE/)+", "/en-AE/", parsed.path or "/")
    path = re.sub(r"/+", "/", path)
    if not path.startswith("/"):
        path = f"/{path}"

    normalized = f"{scheme}://{netloc}{path}"
    query = parsed.query
    if query:
        normalized = f"{normalized}?{query}"
    return normalized.rstrip("/")


def random_delay() -> float:
    return random.uniform(*RANDOM_SLEEP)


def parse_price(value: str) -> Optional[float]:
    if not value:
        return None
    cleaned = re.sub(r"[^0-9.,]", "", value)
    cleaned = cleaned.replace(",", "")
    try:
        return float(cleaned)
    except ValueError:
        return None


def looks_like_service(url: str, context: str = "") -> bool:
    if not url.startswith(BASE_URL):
        return False
    if EXCLUDE_PATTERNS.search(url):
        return False
    if SERVICE_KEYWORDS.search(url):
        return True
    if context and SERVICE_KEYWORDS.search(context):
        return True
    return False


def extract_city_from_url(url: str) -> Optional[str]:
    parsed = urlparse(url)
    segments = [seg for seg in parsed.path.split("/") if seg]
    for city in CITY_PRIORITY:
        if city in segments:
            return city
    return None


def short_uuid(value: str) -> str:
    import uuid

    return str(uuid.uuid5(uuid.NAMESPACE_URL, value))


###############################################################################
# Core scraper                                                                 #
###############################################################################


class JustlifeScraper:
    def __init__(self) -> None:
        self.pages: List[ServiceRecord] = []
        self.options: List[OptionRecord] = []
        self.prices: List[PriceRecord] = []
        self.addons: List[AddonRecord] = []
        self.availability: List[AvailabilityRecord] = []
        self.offers: List[OfferRecord] = []
        self.elements: List[ElementRecord] = []
        self.errors: List[ErrorRecord] = []
        self.discovery_sources: Dict[str, int] = defaultdict(int)
        self.summary_stats: Dict[str, int] = defaultdict(int)
        self.price_signature: Dict[Tuple[str, str], PriceRecord] = {}
        self.combos_attempted: Dict[str, int] = defaultdict(int)

    # ------------------------------------------------------------------
    def log(self, message: str) -> None:
        print(message)

    # ------------------------------------------------------------------
    async def discover_urls(self, browser: Browser) -> List[ServiceCandidate]:
        candidates: Dict[str, ServiceCandidate] = {}

        def register(url: str, source: str, context: str = "") -> None:
            canonical = canonicalize_url(url)
            if not looks_like_service(canonical, context):
                return
            if canonical not in candidates:
                candidates[canonical] = ServiceCandidate(
                    url=canonical,
                    source=source,
                    is_checkout="/checkout" in canonical,
                    city=extract_city_from_url(canonical),
                )
                self.discovery_sources[source] += 1

        context = await browser.new_context(
            user_agent=USER_AGENT,
            locale="en-AE",
            extra_http_headers={"accept-language": "en-AE"},
        )
        context.set_default_timeout(15000)
        page = await context.new_page()
        try:
            await page.goto(BASE_URL, wait_until="networkidle", timeout=60000)
            await page.wait_for_timeout(1200)
            await self.ensure_location(page)
            await self.close_modals(page)

            async def collect(selector: str, source: str) -> None:
                try:
                    anchors: List[Dict[str, Any]] = await page.eval_on_selector_all(
                        selector,
                        "els => els.map(el => ({\n                            href: el.href || '',\n                            text: (el.innerText || '').trim(),\n                            aria: el.getAttribute('aria-label') || '',\n                            classes: el.className || ''\n                        }))",
                    )
                except Exception:
                    return
                for anchor in anchors:
                    href = anchor.get("href") or ""
                    if not href:
                        continue
                    context_text = " ".join(
                        filter(
                            None,
                            [
                                (anchor.get("text") or "").lower(),
                                (anchor.get("aria") or "").lower(),
                                (anchor.get("classes") or "").lower(),
                            ],
                        )
                    )
                    register(href, source, context_text)

            await collect("nav a[href]", "nav")
            await collect("footer a[href]", "footer")
            await collect("section a[href]", "section")
            await collect("a[href*='/en-AE/']", "home")
        finally:
            await page.close()
            await context.close()

        def fetch(url: str) -> Optional[str]:
            try:
                resp = requests.get(url, timeout=30, headers={"User-Agent": USER_AGENT, "accept-language": "en-AE"})
                if resp.ok:
                    return resp.text
            except Exception:
                return None
            return None

        sitemap_html = fetch(f"{BASE_URL}/sitemap.xml")
        if sitemap_html and "<urlset" in sitemap_html:
            soup = BeautifulSoup(sitemap_html, "xml")
            for loc in soup.find_all("loc"):
                register(loc.text.strip(), "sitemap", "sitemap")

        if len(candidates) < 50:
            self.log(
                "[WARN] Discovery found fewer than 50 candidates; verify filters and widget availability."
            )

        discovered = list(candidates.values())
        self.log(f"Discovery collected {len(discovered)} candidate URLs")
        return discovered

    # ------------------------------------------------------------------
    async def ensure_location(self, page: Page) -> None:
        try:
            chip = page.locator("[data-testid*='location']").first
            if await chip.count() == 0:
                chip = page.locator("text=/Dubai/i").first
            if await chip.count() > 0:
                text = (await chip.inner_text()).lower()
                if "dubai" not in text:
                    await chip.click()
                    await page.wait_for_timeout(600)
                    option = page.locator("text=/Dubai/i").first
                    if await option.count() > 0:
                        await option.click()
                        await page.wait_for_load_state("networkidle")
        except PlaywrightTimeout:
            self.log("[WARN] Location chip interaction timed out")
        except Exception as exc:
            self.log(f"[WARN] Location handling error: {exc}")

    # ------------------------------------------------------------------
    async def close_modals(self, page: Page) -> None:
        modal_selectors = [
            "button:has-text('Skip')",
            "button:has-text('Close')",
            "button[aria-label='Close']",
            "[data-testid*='close']",
        ]
        for selector in modal_selectors:
            locator = page.locator(selector)
            if await locator.count() > 0:
                with contextlib.suppress(Exception):
                    await locator.first.click()
                    await page.wait_for_timeout(400)

    # ------------------------------------------------------------------
    async def find_booking_widget(self, page: Page) -> bool:
        for selector in BOOKING_WIDGET_PROBES:
            locator = page.locator(selector)
            if await locator.count() > 0:
                return True
        return False

    # ------------------------------------------------------------------
    async def reach_booking_flow(self, page: Page, candidate: ServiceCandidate) -> Tuple[bool, str]:
        await self.ensure_location(page)
        await self.close_modals(page)
        if await self.find_booking_widget(page):
            return True, page.url

        # Look for anchor with checkout href
        anchors = page.locator("a[href*='/checkout']")
        if await anchors.count() > 0:
            href = await anchors.first.get_attribute("href")
            if href:
                target = urljoin(page.url, href)
                await page.goto(target, wait_until="networkidle")
                await self.close_modals(page)
                await page.wait_for_timeout(600)
                if await self.find_booking_widget(page):
                    return True, page.url

        # Attempt CTA clicks
        buttons = page.locator("button, a")
        for pattern in CTA_TEXT_PATTERNS:
            locator = buttons.filter(has_text=re.compile(pattern, re.IGNORECASE))
            if await locator.count() > 0:
                with contextlib.suppress(Exception):
                    await locator.first.click()
                    await page.wait_for_load_state("networkidle")
                    await page.wait_for_timeout(800)
                    await self.close_modals(page)
                    if await self.find_booking_widget(page):
                        return True, page.url
        return False, page.url

    # ------------------------------------------------------------------
    async def extract_service_name(self, page: Page) -> str:
        for selector in ["h1", "h2", "[data-testid*='title']"]:
            locator = page.locator(selector).first
            if await locator.count() > 0:
                text = (await locator.inner_text()).strip()
                if text:
                    return re.sub(r"\s+", " ", text)
        title = await page.title()
        return title.strip() or "Service"

    # ------------------------------------------------------------------
    async def screenshot(self, page: Page, name: str) -> str:
        filename = f"{slugify(name)}.png"
        path = SHOTS_DIR / filename
        await page.screenshot(path=path, full_page=True)
        return str(path.relative_to(OUTPUT_DIR))

    # ------------------------------------------------------------------
    def record_element(self, page_id: str, selector: str, role: str, text: str) -> None:
        cleaned = re.sub(r"\s+", " ", (text or "").strip())
        self.elements.append(ElementRecord(page_id, selector, role, text or "", cleaned))

    # ------------------------------------------------------------------
    async def watch_price_change(self, page: Page, action) -> Tuple[Optional[str], Optional[str]]:
        before_text: Optional[str] = None
        for selector in PRICE_SELECTORS:
            locator = page.locator(selector).first
            if await locator.count() > 0:
                with contextlib.suppress(Exception):
                    before_text = await locator.inner_text()
                    break
        await action()
        await page.wait_for_timeout(350)
        timeout_at = time.monotonic() + 8
        after_text: Optional[str] = None
        while time.monotonic() < timeout_at:
            for selector in PRICE_SELECTORS:
                locator = page.locator(selector).first
                if await locator.count() > 0:
                    with contextlib.suppress(Exception):
                        candidate = await locator.inner_text()
                        if candidate and candidate != before_text:
                            after_text = candidate
                            return before_text, after_text
            await page.wait_for_timeout(300)
        return before_text, after_text

    # ------------------------------------------------------------------
    async def capture_totals(self, page: Page, page_id: str) -> Tuple[Optional[float], Optional[float], Optional[float], Optional[float], Optional[str], str]:
        summary_text = []
        grand_total = base_price = vat_amount = fees_total = None
        vat_percent = None
        currency = "AED"
        for selector in PRICE_SELECTORS:
            locator = page.locator(selector)
            if await locator.count() > 0:
                with contextlib.suppress(Exception):
                    text = await locator.first.inner_text()
                    if text:
                        summary_text.append(text)
                        candidate = parse_price(text)
                        if candidate:
                            if not grand_total or candidate > grand_total:
                                grand_total = candidate
        summary_locator = page.locator(
            " | ".join(f"[class*='{hint}']" for hint in SUMMARY_HINTS)
        )
        if await summary_locator.count() > 0:
            with contextlib.suppress(Exception):
                summary_text.append(await summary_locator.first.inner_text())
        merged = "\n".join(summary_text)
        if merged:
            self.record_element(page_id, "summary", "price_summary", merged)
            vat_match = re.search(r"VAT\s*(\d+\.?\d*)%", merged, re.IGNORECASE)
            if vat_match:
                vat_percent = float(vat_match.group(1))
            vat_amt_match = re.search(r"VAT[^0-9]*([0-9,\.]+)", merged, re.IGNORECASE)
            if vat_amt_match:
                vat_amount = parse_price(vat_amt_match.group(1))
            base_match = re.search(r"Base[^0-9]*([0-9,\.]+)", merged, re.IGNORECASE)
            if base_match:
                base_price = parse_price(base_match.group(1))
            fees_match = re.search(r"Fee[^0-9]*([0-9,\.]+)", merged, re.IGNORECASE)
            if fees_match:
                fees_total = parse_price(fees_match.group(1))
        return base_price, fees_total, vat_amount, vat_percent, currency, merged

    # ------------------------------------------------------------------
    async def capture_addons(self, page: Page, page_id: str, service_name: str, service_url: str) -> bool:
        addon_headers = page.locator("text=/Add[- ]?ons|Extras|Upgrades/i")
        captured = False
        if await addon_headers.count() == 0:
            return captured
        header = addon_headers.first
        await header.scroll_into_view_if_needed()
        container = header.locator("xpath=ancestor::section[1]")
        if await container.count() == 0:
            container = header
        cards = container.locator("xpath=.//div[contains(@class,'card') or contains(@class,'addon')]")
        if await cards.count() == 0:
            cards = container.locator("xpath=.//li")
        for idx in range(await cards.count()):
            card = cards.nth(idx)
            with contextlib.suppress(Exception):
                text = await card.inner_text()
                name = re.split(r"\n", text.strip())[0]
                price = parse_price(text)
                recommended = "recommended" in text.lower() or "popular" in text.lower()
                required = "required" in text.lower() or "mandatory" in text.lower()
                self.addons.append(
                    AddonRecord(
                        page_id=page_id,
                        addon_name=name,
                        addon_description=text.strip(),
                        addon_price=price,
                        currency="AED" if price else None,
                        is_recommended=recommended,
                        is_required=required,
                        seen_on_step="checkout",
                        pricing_source_url=service_url,
                        collected_at=utc_now(),
                    )
                )
                captured = True
        return captured

    # ------------------------------------------------------------------
    async def capture_availability(self, page: Page, page_id: str, service_url: str) -> bool:
        calendar_locator = page.locator("[data-testid*='date'], .calendar, input[type='date']")
        if await calendar_locator.count() == 0:
            return False
        calendar = calendar_locator.first
        with contextlib.suppress(Exception):
            await calendar.scroll_into_view_if_needed()
            await calendar.click()
            await page.wait_for_timeout(600)
        captured = False
        today = datetime.now(timezone.utc).date()
        for _ in range(DAYS_AHEAD):
            slots = page.locator("[data-testid*='slot'], .time-slot, .slot")
            count = await slots.count()
            if count == 0:
                break
            for i in range(count):
                slot = slots.nth(i)
                with contextlib.suppress(Exception):
                    label = (await slot.inner_text()).strip()
                    if not label:
                        continue
                    status = "available"
                    classes = await slot.get_attribute("class") or ""
                    if "disabled" in classes or "unavailable" in classes:
                        status = "unavailable"
                    date_attr = await slot.get_attribute("data-date")
                    if date_attr:
                        date_obj = datetime.fromisoformat(date_attr)
                        date_str = date_obj.date().isoformat()
                        lead_time = (date_obj.date() - today).days
                    else:
                        date_str = (today + timedelta(days=_)).isoformat()
                        lead_time = _
                    self.availability.append(
                        AvailabilityRecord(
                            page_id=page_id,
                            date_str=date_str,
                            time_slot_label=label,
                            slot_status=status,
                            lead_time_days=lead_time,
                            collected_at=utc_now(),
                        )
                    )
                    captured = True
            next_button = page.locator("button:has-text('Next')")
            if await next_button.count() > 0:
                with contextlib.suppress(Exception):
                    await next_button.first.click()
                    await page.wait_for_timeout(500)
            else:
                break
        return captured

    # ------------------------------------------------------------------
    def compute_model(self, control_names: List[str], card_sections: List[str]) -> str:
        for model, hints in CONTROL_LABEL_HINTS.items():
            for name in control_names + card_sections:
                if any(hint in name.lower() for hint in hints):
                    return model
        return "OTHER"

    # ------------------------------------------------------------------
    async def enumerate_sections(self, page: Page, page_id: str, service_name: str, service_url: str) -> Tuple[List[str], int]:
        section_titles: List[str] = []
        price_rows = 0
        sections = page.locator(
            "xpath=//section[.//button[contains(translate(text(),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'add')]]"
        )
        sections_count = await sections.count()
        for idx in range(sections_count):
            if self.combos_attempted[page_id] >= MAX_COMBOS_PER_SERVICE:
                break
            section = sections.nth(idx)
            try:
                title_locator = section.locator("xpath=.//h2 | .//h3 | .//h4 | .//p")
                section_name = f"Section {idx+1}"
                if await title_locator.count() > 0:
                    section_name = (await title_locator.first.inner_text()).strip() or section_name
                section_titles.append(section_name)
                cards = section.locator(
                    "xpath=.//div[contains(@class,'card') or contains(@class,'tile') or contains(@data-testid,'card')]"
                )
                if await cards.count() == 0:
                    cards = section.locator("xpath=.//li")
                for card_index in range(min(await cards.count(), 10)):
                    if self.combos_attempted[page_id] >= MAX_COMBOS_PER_SERVICE:
                        break
                    card = cards.nth(card_index)
                    await card.scroll_into_view_if_needed()
                    text = (await card.inner_text()).strip()
                    header = re.split(r"\n", text)[0]
                    quantity_buttons = card.locator("button:has-text('Add')")
                    if await quantity_buttons.count() == 0:
                        quantity_buttons = card.locator("button:has-text('+')")
                    if await quantity_buttons.count() == 0:
                        continue
                    option_code = slugify(f"{section_name}-{header}")
                    self.options.append(
                        OptionRecord(
                            page_id=page_id,
                            control_group=section_name,
                            option_code=option_code,
                            option_label=header,
                            option_type="card",
                            is_default=False,
                            sort_index=card_index,
                        )
                    )

                    async def click_add() -> None:
                        await quantity_buttons.first.click()
                    before, after = await self.watch_price_change(page, click_add)
                    base_price, fees_total, vat_amount, vat_percent, currency, summary = await self.capture_totals(page, page_id)
                    qa_flags = []
                    if not after:
                        qa_flags.append("NO_PRICE_DELTA")
                    record = PriceRecord(
                        page_id=page_id,
                        service_name=service_name,
                        model_type="FLAT_PACKAGE",
                        hours=None,
                        pros=None,
                        bedrooms=None,
                        bathrooms=None,
                        kitchen_package=None,
                        appliance_combo=header,
                        package_name=section_name,
                        duration=None,
                        offer_state="OFF",
                        offer_name=None,
                        discount_type=None,
                        discount_value=None,
                        base_price=base_price,
                        fees_total=fees_total,
                        vat_amount=vat_amount,
                        vat_percent=vat_percent,
                        grand_total=parse_price(after or before or summary),
                        currency=currency,
                        pricing_source_url=service_url,
                        collected_at=utc_now(),
                        qa_flags=",".join(qa_flags),
                        screenshot_ref=None,
                    )
                    signature = (page_id, f"package:{section_name}:{header}")
                    if signature not in self.price_signature:
                        self.prices.append(record)
                        self.price_signature[signature] = record
                        price_rows += 1
                        self.combos_attempted[page_id] += 1
                    # Attempt to remove item to reset state
                    remove_btn = page.locator("button:has-text('Remove')")
                    if await remove_btn.count() > 0:
                        with contextlib.suppress(Exception):
                            await remove_btn.first.click()
                            await page.wait_for_timeout(400)
            except Exception as exc:
                self.errors.append(ErrorRecord(page_id, service_url, "section-enumeration", str(exc)))
        return section_titles, price_rows

    # ------------------------------------------------------------------
    async def enumerate_hours_pros(self, page: Page, page_id: str, service_name: str, service_url: str) -> Tuple[List[str], int]:
        titles: List[str] = []
        price_rows = 0
        hours_options: List[str] = []
        pros_options: List[str] = []

        for selector in ["select[name*='hour']", "select[data-testid*='hour']"]:
            locator = page.locator(selector)
            if await locator.count() > 0:
                try:
                    option_data = await page.eval_on_selector(
                        selector,
                        "el => Array.from(el.options).map(opt => ({value: opt.value, label: (opt.label || opt.textContent || '').trim()}))",
                    )
                except Exception:
                    option_data = []
                hours_options = [opt["label"] or opt["value"] for opt in option_data if opt.get("value")]
                for idx, opt in enumerate(option_data):
                    self.options.append(
                        OptionRecord(
                            page_id=page_id,
                            control_group="Hours",
                            option_code=opt.get("value") or opt.get("label") or f"hours-{idx}",
                            option_label=opt.get("label") or opt.get("value") or "",
                            option_type="select",
                            is_default=idx == 0,
                            sort_index=idx,
                        )
                    )
                titles.append("Hours")
        if not hours_options:
            button_groups = page.locator("[data-testid*='hour'] button, [data-testid*='duration'] button, button:has-text('hour')")
            count = await button_groups.count()
            for idx in range(min(count, 6)):
                label = (await button_groups.nth(idx).inner_text()).strip()
                hours_options.append(label)
                self.options.append(
                    OptionRecord(
                        page_id=page_id,
                        control_group="Hours",
                        option_code=slugify(f"hours-{label}") or f"hours-{idx}",
                        option_label=label,
                        option_type="button",
                        is_default=idx == 0,
                        sort_index=idx,
                    )
                )
            if hours_options:
                titles.append("Hours")

        for selector in ["select[name*='pro']", "select[data-testid*='pro']", "select[name*='cleaner']"]:
            locator = page.locator(selector)
            if await locator.count() > 0:
                try:
                    option_data = await page.eval_on_selector(
                        selector,
                        "el => Array.from(el.options).map(opt => ({value: opt.value, label: (opt.label || opt.textContent || '').trim()}))",
                    )
                except Exception:
                    option_data = []
                pros_options = [opt["label"] or opt["value"] for opt in option_data if opt.get("value")]
                for idx, opt in enumerate(option_data):
                    self.options.append(
                        OptionRecord(
                            page_id=page_id,
                            control_group="Professionals",
                            option_code=opt.get("value") or opt.get("label") or f"pros-{idx}",
                            option_label=opt.get("label") or opt.get("value") or "",
                            option_type="select",
                            is_default=idx == 0,
                            sort_index=idx,
                        )
                    )
                titles.append("Pros")
        if not pros_options:
            button_groups = page.locator("[data-testid*='pro'] button, button:has-text('Cleaner'), button:has-text('cleaner')")
            count = await button_groups.count()
            for idx in range(min(count, 4)):
                label = (await button_groups.nth(idx).inner_text()).strip()
                pros_options.append(label)
                self.options.append(
                    OptionRecord(
                        page_id=page_id,
                        control_group="Professionals",
                        option_code=slugify(f"pros-{label}") or f"pros-{idx}",
                        option_label=label,
                        option_type="button",
                        is_default=idx == 0,
                        sort_index=idx,
                    )
                )
            if pros_options:
                titles.append("Pros")

        if not hours_options or not pros_options:
            return titles, price_rows

        hours_sample = hours_options[:3]
        pros_sample = pros_options[:3]
        for hour in hours_sample:
            for pro in pros_sample:
                if self.combos_attempted[page_id] >= MAX_COMBOS_PER_SERVICE:
                    break

                async def set_options() -> None:
                    for selector in ["select[name*='hour']", "select[data-testid*='hour']"]:
                        sel = page.locator(selector)
                        if await sel.count() > 0:
                            with contextlib.suppress(Exception):
                                await sel.select_option(label=hour)
                    for selector in ["select[name*='pro']", "select[data-testid*='pro']", "select[name*='cleaner']"]:
                        sel = page.locator(selector)
                        if await sel.count() > 0:
                            with contextlib.suppress(Exception):
                                await sel.select_option(label=pro)

                before, after = await self.watch_price_change(page, set_options)
                base_price, fees_total, vat_amount, vat_percent, currency, summary = await self.capture_totals(page, page_id)
                qa_flags = []
                if not after and not before:
                    qa_flags.append("NO_PRICE")
                signature = (page_id, f"hours:{hour}|pros:{pro}")
                if signature not in self.price_signature:
                    self.prices.append(
                        PriceRecord(
                            page_id=page_id,
                            service_name=service_name,
                            model_type="HOURS_PROS",
                            hours=hour,
                            pros=pro,
                            bedrooms=None,
                            bathrooms=None,
                            kitchen_package=None,
                            appliance_combo=None,
                            package_name=None,
                            duration=None,
                            offer_state="OFF",
                            offer_name=None,
                            discount_type=None,
                            discount_value=None,
                            base_price=base_price,
                            fees_total=fees_total,
                            vat_amount=vat_amount,
                            vat_percent=vat_percent,
                            grand_total=parse_price(after or before or summary),
                            currency=currency,
                            pricing_source_url=service_url,
                            collected_at=utc_now(),
                            qa_flags=",".join(qa_flags),
                            screenshot_ref=None,
                        )
                    )
                    price_rows += 1
                    self.combos_attempted[page_id] += 1
        return titles, price_rows

    # ------------------------------------------------------------------
    async def process_service(self, context: BrowserContext, candidate: ServiceCandidate) -> None:
        page = await context.new_page()
        page.set_default_timeout(15000)
        page_id = short_uuid(candidate.url)
        try:
            await page.goto(candidate.url, wait_until="networkidle", timeout=60000)
            await page.wait_for_timeout(800)
            widget_ready, final_url = await self.reach_booking_flow(page, candidate)
            screenshot_path = await self.screenshot(page, f"service_{slugify(candidate.url)}")
            if not widget_ready:
                self.errors.append(ErrorRecord(page_id, candidate.url, "no-widget", "Booking widget not detected", screenshot_path))
                self.pages.append(
                    ServiceRecord(
                        page_id=page_id,
                        parent_page_id=None,
                        level="service",
                        url=candidate.url,
                        page_title_text=await self.extract_service_name(page),
                        detected_model_type="OTHER",
                        has_addons=False,
                        has_slots=False,
                        notes="Widget not detected",
                    )
                )
                return

            service_name = await self.extract_service_name(page)
            control_labels: List[str] = []
            total_price_rows = 0

            titles_hours, rows_hours = await self.enumerate_hours_pros(page, page_id, service_name, final_url)
            control_labels.extend(titles_hours)
            total_price_rows += rows_hours

            section_titles, rows_sections = await self.enumerate_sections(page, page_id, service_name, final_url)
            control_labels.extend(section_titles)
            total_price_rows += rows_sections

            addons_captured = await self.capture_addons(page, page_id, service_name, final_url)
            slots_captured = await self.capture_availability(page, page_id, final_url)

            model_type = self.compute_model(control_labels, section_titles)
            self.pages.append(
                ServiceRecord(
                    page_id=page_id,
                    parent_page_id=None,
                    level="service",
                    url=final_url,
                    page_title_text=service_name,
                    detected_model_type=model_type,
                    has_addons=addons_captured,
                    has_slots=slots_captured,
                    notes=f"controls={control_labels}, sections={section_titles}, combos={total_price_rows}",
                )
            )
            if total_price_rows == 0:
                self.errors.append(ErrorRecord(page_id, final_url, "pricing", "No price combinations captured", screenshot_path))
        except Exception as exc:
            screenshot_path = None
            with contextlib.suppress(Exception):
                screenshot_path = await self.screenshot(page, f"error_{slugify(candidate.url)}")
            self.errors.append(ErrorRecord(page_id, candidate.url, "runtime", str(exc), screenshot_path))
        finally:
            await page.close()

    # ------------------------------------------------------------------
    async def run(self) -> None:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=HEADLESS)
            candidates = await self.discover_urls(browser)
            kept = candidates[:500]
            context = await browser.new_context(
                user_agent=USER_AGENT,
                locale="en-AE",
                extra_http_headers={"accept-language": "en-AE"},
            )
            context.set_default_timeout(15000)
            try:
                for candidate in kept:
                    await self.process_service(context, candidate)
                    await asyncio.sleep(random_delay())
            finally:
                await context.close()
                await browser.close()
        self.summary_stats["services"] = len(self.pages)
        self.summary_stats["prices"] = len(self.prices)
        self.summary_stats["addons"] = len(self.addons)
        self.summary_stats["availability"] = len(self.availability)
        self.summary_stats["offers"] = len(self.offers)

    # ------------------------------------------------------------------
    def to_dataframe(self, records: List[Any], columns: List[str]) -> pd.DataFrame:
        return pd.DataFrame([asdict(record) for record in records], columns=columns)

    # ------------------------------------------------------------------
    def export_outputs(self) -> None:
        pages_df = self.to_dataframe(
            self.pages,
            [
                "page_id",
                "parent_page_id",
                "level",
                "url",
                "page_title_text",
                "detected_model_type",
                "has_addons",
                "has_slots",
                "notes",
            ],
        )
        options_df = self.to_dataframe(
            self.options,
            [
                "page_id",
                "control_group",
                "option_code",
                "option_label",
                "option_type",
                "is_default",
                "sort_index",
            ],
        )
        prices_df = self.to_dataframe(
            self.prices,
            [
                "page_id",
                "service_name",
                "model_type",
                "hours",
                "pros",
                "bedrooms",
                "bathrooms",
                "kitchen_package",
                "appliance_combo",
                "package_name",
                "duration",
                "offer_state",
                "offer_name",
                "discount_type",
                "discount_value",
                "base_price",
                "fees_total",
                "vat_amount",
                "vat_percent",
                "grand_total",
                "currency",
                "pricing_source_url",
                "collected_at",
                "qa_flags",
                "screenshot_ref",
            ],
        )
        addons_df = self.to_dataframe(
            self.addons,
            [
                "page_id",
                "addon_name",
                "addon_description",
                "addon_price",
                "currency",
                "is_recommended",
                "is_required",
                "seen_on_step",
                "pricing_source_url",
                "collected_at",
            ],
        )
        availability_df = self.to_dataframe(
            self.availability,
            [
                "page_id",
                "date_str",
                "time_slot_label",
                "slot_status",
                "lead_time_days",
                "collected_at",
            ],
        )
        offers_df = self.to_dataframe(
            self.offers,
            [
                "page_id",
                "offer_name",
                "offer_badge_text",
                "description",
                "eligibility",
                "discount_type",
                "discount_value",
                "stacking_rules",
                "collected_at",
            ],
        )
        elements_df = self.to_dataframe(
            self.elements,
            [
                "page_id",
                "selector",
                "element_role",
                "raw_text",
                "normalized_text",
            ],
        )
        errors_df = self.to_dataframe(
            self.errors,
            [
                "page_id",
                "url",
                "context",
                "message",
                "screenshot",
            ],
        )

        pages_df.to_csv(OUTPUT_DIR / "pages.csv", index=False)
        options_df.to_csv(OUTPUT_DIR / "options.csv", index=False)
        prices_df.to_csv(OUTPUT_DIR / "prices_raw.csv", index=False)
        addons_df.to_csv(OUTPUT_DIR / "addons.csv", index=False)
        availability_df.to_csv(OUTPUT_DIR / "availability.csv", index=False)
        offers_df.to_csv(OUTPUT_DIR / "offers.csv", index=False)
        elements_df.to_csv(OUTPUT_DIR / "elements.csv", index=False)

        notes = OUTPUT_DIR / "NOTES.txt"
        with notes.open("w", encoding="utf-8") as handle:
            handle.write("Justlife scrape run notes\n")
            handle.write(f"Timestamp: {utc_now()}\n")
            handle.write(f"Discovered sources: {dict(self.discovery_sources)}\n")
            handle.write(f"Pages captured: {len(self.pages)}\n")
            handle.write(f"Prices captured: {len(self.prices)}\n")
            if not self.prices:
                handle.write("WARNING: No price rows captured. Investigate selectors and widget flows.\n")
            for error in self.errors[:20]:
                handle.write(f"Error[{error.context}] {error.url}: {error.message}\n")

        self.build_excel(
            pages_df,
            options_df,
            prices_df,
            addons_df,
            availability_df,
            offers_df,
            elements_df,
            errors_df,
        )

    # ------------------------------------------------------------------
    def build_excel(
        self,
        pages: pd.DataFrame,
        options: pd.DataFrame,
        prices: pd.DataFrame,
        addons: pd.DataFrame,
        availability: pd.DataFrame,
        offers: pd.DataFrame,
        elements: pd.DataFrame,
        errors: pd.DataFrame,
    ) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        run_timestamp = utc_now()

        def add_sheet(name: str, df: pd.DataFrame, freeze: bool = True, currency_cols: Optional[List[str]] = None) -> None:
            ws = wb.create_sheet(name[:31])
            data = df.copy()
            if data.empty and data.columns.size == 0:
                data = pd.DataFrame({"Notice": ["No data captured"]})
            for col_idx, column_name in enumerate(data.columns, 1):
                header_cell = ws.cell(row=1, column=col_idx, value=column_name)
                header_cell.font = Font(bold=True)
                header_cell.fill = PatternFill("solid", fgColor="DDEBF7")
                header_cell.alignment = Alignment(horizontal="center")
            for row_idx, (_, row) in enumerate(data.iterrows(), start=2):
                for col_idx, value in enumerate(row, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            if freeze:
                ws.freeze_panes = "A2"
            for col_idx in range(1, len(data.columns) + 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = 20
            if currency_cols:
                for column in currency_cols:
                    if column in data.columns:
                        idx = list(data.columns).index(column) + 1
                        for row_idx in range(2, len(data) + 2):
                            cell = ws.cell(row=row_idx, column=idx)
                            cell.number_format = "#,##0.00"
            ws.auto_filter.ref = ws.dimensions

        def first_non_empty(series: pd.Series) -> Optional[str]:
            for value in series:
                if pd.notna(value) and str(value).strip():
                    return str(value).strip()
            return None

        services_sheet = pd.DataFrame(
            columns=[
                "Service Name",
                "City",
                "Category",
                "Checkout URL",
                "Description",
                "Duration",
                "Base Price",
                "Price AED",
                "Discount",
                "Currency",
                "Add-ons",
                "Combinations JSON",
                "Available Slots",
                "Special Offers",
                "Scrape Timestamp",
            ]
        )

        if not pages.empty:
            services_base = pages.copy()
            services_base["city"] = services_base["url"].apply(lambda u: extract_city_from_url(u) or "")

            if not prices.empty:
                price_stats = (
                    prices.groupby("page_id").agg(
                        base_price_min=("base_price", "min"),
                        grand_total_min=("grand_total", "min"),
                        grand_total_avg=("grand_total", "mean"),
                    )
                ).reset_index()
                duration_map = (
                    prices.groupby("page_id")["duration"].apply(first_non_empty).reset_index(name="duration")
                )
                discount_map = (
                    prices.groupby("page_id")["discount_value"].apply(first_non_empty).reset_index(name="discount_value")
                )
                currency_map = (
                    prices.groupby("page_id")["currency"].apply(first_non_empty).reset_index(name="currency")
                )
                def sanitize(value: Any) -> Any:
                    return None if pd.isna(value) else value

                combos_json = (
                    prices.groupby("page_id")
                    .apply(
                        lambda df: json.dumps(
                            [
                                {
                                    "model_type": sanitize(row.get("model_type")),
                                    "hours": sanitize(row.get("hours")),
                                    "pros": sanitize(row.get("pros")),
                                    "bedrooms": sanitize(row.get("bedrooms")),
                                    "bathrooms": sanitize(row.get("bathrooms")),
                                    "kitchen_package": sanitize(row.get("kitchen_package")),
                                    "appliance_combo": sanitize(row.get("appliance_combo")),
                                    "package_name": sanitize(row.get("package_name")),
                                    "duration": sanitize(row.get("duration")),
                                    "grand_total": sanitize(row.get("grand_total")),
                                }
                                for _, row in df.iterrows()
                            ],
                            ensure_ascii=False,
                        )
                    )
                    .reset_index(name="combinations_json")
                )
                combos_count = (
                    prices.groupby("page_id").size().reset_index(name="combination_count")
                )
            else:
                price_stats = pd.DataFrame(columns=["page_id", "base_price_min", "grand_total_min", "grand_total_avg"])
                duration_map = pd.DataFrame(columns=["page_id", "duration"])
                discount_map = pd.DataFrame(columns=["page_id", "discount_value"])
                currency_map = pd.DataFrame(columns=["page_id", "currency"])
                combos_json = pd.DataFrame(columns=["page_id", "combinations_json"])
                combos_count = pd.DataFrame(columns=["page_id", "combination_count"])

            if not addons.empty:
                addons_group = (
                    addons.groupby("page_id")["addon_name"]
                    .apply(lambda s: ", ".join(sorted({name.strip() for name in s if isinstance(name, str) and name.strip()})))
                    .reset_index(name="addon_names")
                )
            else:
                addons_group = pd.DataFrame(columns=["page_id", "addon_names"])

            if not availability.empty:
                availability_group = (
                    availability.groupby("page_id")["time_slot_label"].count().reset_index(name="available_slots")
                )
            else:
                availability_group = pd.DataFrame(columns=["page_id", "available_slots"])

            if not offers.empty:
                offers_group = (
                    offers.groupby("page_id")["offer_name"]
                    .apply(lambda s: ", ".join(sorted({name.strip() for name in s if isinstance(name, str) and name.strip()})))
                    .reset_index(name="offer_names")
                )
            else:
                offers_group = pd.DataFrame(columns=["page_id", "offer_names"])

            services_enriched = (
                services_base.merge(price_stats, on="page_id", how="left")
                .merge(duration_map, on="page_id", how="left")
                .merge(discount_map, on="page_id", how="left")
                .merge(currency_map, on="page_id", how="left")
                .merge(combos_json, on="page_id", how="left")
                .merge(combos_count, on="page_id", how="left")
                .merge(addons_group, on="page_id", how="left")
                .merge(availability_group, on="page_id", how="left")
                .merge(offers_group, on="page_id", how="left")
            )

            services_enriched["Base Price"] = services_enriched.get("base_price_min")
            services_enriched["Price AED"] = services_enriched.get("grand_total_min")
            services_enriched.loc[
                services_enriched["Base Price"].isna(), "Base Price"
            ] = services_enriched["Price AED"]
            services_enriched["Currency"] = services_enriched.get("currency")
            services_enriched.loc[
                services_enriched["Currency"].isna() & services_enriched["Price AED"].notna(), "Currency"
            ] = "AED"

            available_slots_series = services_enriched.get("available_slots")
            if available_slots_series is None:
                available_slots_series = pd.Series([0] * len(services_enriched))
            available_slots_series = pd.to_numeric(available_slots_series, errors="coerce").fillna(0).astype(int)

            services_sheet = services_enriched.assign(
                **{
                    "Service Name": services_enriched["page_title_text"],
                    "City": services_enriched["city"],
                    "Category": services_enriched["detected_model_type"],
                    "Checkout URL": services_enriched["url"],
                    "Description": services_enriched["notes"],
                    "Duration": services_enriched.get("duration"),
                    "Discount": services_enriched.get("discount_value"),
                    "Add-ons": services_enriched.get("addon_names"),
                    "Combinations JSON": services_enriched.get("combinations_json"),
                    "Available Slots": available_slots_series,
                    "Special Offers": services_enriched.get("offer_names"),
                    "Scrape Timestamp": run_timestamp,
                }
            )[
                [
                    "Service Name",
                    "City",
                    "Category",
                    "Checkout URL",
                    "Description",
                    "Duration",
                    "Base Price",
                    "Price AED",
                    "Discount",
                    "Currency",
                    "Add-ons",
                    "Combinations JSON",
                    "Available Slots",
                    "Special Offers",
                    "Scrape Timestamp",
                ]
            ]

            for column in ["Base Price", "Price AED"]:
                services_sheet[column] = pd.to_numeric(services_sheet[column], errors="coerce")

        combinations_sheet = pd.DataFrame(
            columns=[
                "service_id",
                "service_name",
                "service_url",
                "combination_key",
                "model_type",
                "hours",
                "num_professionals",
                "bedrooms",
                "bathrooms",
                "kitchen_package",
                "appliance_combo",
                "package_name",
                "duration",
                "base_price",
                "fees_total",
                "vat_amount",
                "vat_percent",
                "grand_total",
                "currency",
                "offer_state",
                "offer_name",
                "discount_type",
                "discount_value",
                "qa_flags",
                "collected_at",
                "screenshot_ref",
            ]
        )

        if not prices.empty:
            combinations_enriched = prices.merge(
                pages[["page_id", "page_title_text", "url"]], on="page_id", how="left"
            )

            def build_key(row: pd.Series) -> str:
                parts: List[str] = []
                for label, column in [
                    ("hours", "hours"),
                    ("pros", "pros"),
                    ("bedrooms", "bedrooms"),
                    ("bathrooms", "bathrooms"),
                    ("kitchen", "kitchen_package"),
                    ("combo", "appliance_combo"),
                    ("package", "package_name"),
                    ("duration", "duration"),
                ]:
                    if column in row and pd.notna(row[column]) and str(row[column]).strip():
                        parts.append(f"{label}={row[column]}")
                return " | ".join(parts) if parts else "default"

            combinations_enriched = combinations_enriched.assign(
                service_id=combinations_enriched["page_id"],
                service_name=combinations_enriched["service_name"].fillna(combinations_enriched["page_title_text"]),
                service_url=combinations_enriched["pricing_source_url"],
                combination_key=combinations_enriched.apply(build_key, axis=1),
                num_professionals=combinations_enriched["pros"],
            )

            combinations_sheet = combinations_enriched[
                [
                    "service_id",
                    "service_name",
                    "service_url",
                    "combination_key",
                    "model_type",
                    "hours",
                    "num_professionals",
                    "bedrooms",
                    "bathrooms",
                    "kitchen_package",
                    "appliance_combo",
                    "package_name",
                    "duration",
                    "base_price",
                    "fees_total",
                    "vat_amount",
                    "vat_percent",
                    "grand_total",
                    "currency",
                    "offer_state",
                    "offer_name",
                    "discount_type",
                    "discount_value",
                    "qa_flags",
                    "collected_at",
                    "screenshot_ref",
                ]
            ]

        addons_sheet = pd.DataFrame(
            columns=[
                "service_id",
                "service_name",
                "service_url",
                "addon_name",
                "addon_description",
                "min_qty",
                "max_qty",
                "unit_price",
                "currency",
                "is_recommended",
                "is_required",
                "extracted_utc",
            ]
        )

        if not addons.empty:
            addons_enriched = addons.merge(
                pages[["page_id", "page_title_text", "url"]], on="page_id", how="left"
            )
            addons_sheet = addons_enriched.assign(
                service_id=addons_enriched["page_id"],
                service_name=addons_enriched["page_title_text"],
                service_url=addons_enriched["pricing_source_url"],
                min_qty=None,
                max_qty=None,
                unit_price=addons_enriched["addon_price"],
                extracted_utc=addons_enriched["collected_at"],
            )[
                [
                    "service_id",
                    "service_name",
                    "service_url",
                    "addon_name",
                    "addon_description",
                    "min_qty",
                    "max_qty",
                    "unit_price",
                    "currency",
                    "is_recommended",
                    "is_required",
                    "extracted_utc",
                ]
            ]

        availability_sheet = pd.DataFrame(
            columns=[
                "service_id",
                "service_name",
                "service_url",
                "date",
                "time_label",
                "status",
                "lead_time_days",
                "captured_utc",
            ]
        )

        if not availability.empty:
            availability_enriched = availability.merge(
                pages[["page_id", "page_title_text", "url"]], on="page_id", how="left"
            )
            availability_sheet = availability_enriched.assign(
                service_id=availability_enriched["page_id"],
                service_name=availability_enriched["page_title_text"],
                service_url=availability_enriched["url"],
                date=availability_enriched["date_str"],
                time_label=availability_enriched["time_slot_label"],
                status=availability_enriched["slot_status"],
                captured_utc=availability_enriched["collected_at"],
            )[
                [
                    "service_id",
                    "service_name",
                    "service_url",
                    "date",
                    "time_label",
                    "status",
                    "lead_time_days",
                    "captured_utc",
                ]
            ]

        errors_sheet = pd.DataFrame(
            columns=["service_id", "service_name", "url", "context", "message", "screenshot"]
        )

        if not errors.empty:
            errors_enriched = errors.merge(
                pages[["page_id", "page_title_text"]], on="page_id", how="left"
            )
            errors_sheet = errors_enriched.assign(
                service_id=errors_enriched["page_id"],
                service_name=errors_enriched["page_title_text"],
            )[
                ["service_id", "service_name", "url", "context", "message", "screenshot"]
            ]

        offers_sheet = pd.DataFrame(
            columns=[
                "page_id",
                "offer_name",
                "offer_badge_text",
                "description",
                "eligibility",
                "discount_type",
                "discount_value",
                "stacking_rules",
                "collected_at",
            ]
        )
        if not offers.empty:
            offers_sheet = offers

        selectors_sheet = pd.DataFrame(
            columns=["page_id", "selector", "element_role", "raw_text", "normalized_text"]
        )
        if not elements.empty:
            selectors_sheet = elements

        meta_rows = [
            {"Metric": "Run Timestamp", "Value": run_timestamp},
            {"Metric": "Services processed", "Value": len(pages)},
            {"Metric": "Price combinations", "Value": len(prices)},
            {"Metric": "Add-ons captured", "Value": len(addons)},
            {"Metric": "Availability slots", "Value": len(availability)},
            {"Metric": "Offers captured", "Value": len(offers)},
        ]
        for source, count in sorted(self.discovery_sources.items()):
            meta_rows.append({"Metric": f"Discovery::{source}", "Value": count})
        meta_sheet = pd.DataFrame(meta_rows)

        add_sheet("Services", services_sheet, currency_cols=["Base Price", "Price AED"])
        add_sheet("Combinations", combinations_sheet, currency_cols=["base_price", "fees_total", "vat_amount", "grand_total"])
        add_sheet("AddOns", addons_sheet, currency_cols=["unit_price"])
        add_sheet("Availability", availability_sheet)
        add_sheet("Errors", errors_sheet)
        add_sheet("Meta", meta_sheet, freeze=False)
        add_sheet("Offers", offers_sheet)
        add_sheet("Selectors_Trace", selectors_sheet)

        combo_ws = wb["Combinations"]
        if "grand_total" in combinations_sheet.columns and combinations_sheet.shape[0] > 0:
            grand_col_index = list(combinations_sheet.columns).index("grand_total") + 1
            color_scale = ColorScaleRule(
                start_type="min",
                start_color="F8F9FA",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFF2CC",
                end_type="max",
                end_color="F4B084",
            )
            combo_ws.conditional_formatting.add(
                f"{get_column_letter(grand_col_index)}2:{get_column_letter(grand_col_index)}{combo_ws.max_row}",
                color_scale,
            )
            qa_index = list(combinations_sheet.columns).index("qa_flags") + 1
            qa_rule = CellIsRule(
                operator="notEqual",
                formula=["\"\""],
                fill=PatternFill("solid", fgColor="FFC7CE"),
            )
            combo_ws.conditional_formatting.add(
                f"{get_column_letter(qa_index)}2:{get_column_letter(qa_index)}{combo_ws.max_row}",
                qa_rule,
            )

        workbook_path = OUTPUT_DIR / "Justlife_CRM.xlsx"
        wb.save(workbook_path)

    # ------------------------------------------------------------------
    def print_summary(self) -> None:
        print("Run summary")
        print("============")
        print(f"Discovered URLs: {sum(self.discovery_sources.values())}   Bookable targets: {len(self.pages)}   Skipped: {len(self.errors)}")
        print(
            f"Price combinations: {len(self.prices)}  Add-ons: {len(self.addons)}  Availability slots: {len(self.availability)}  Offers: {len(self.offers)}"
        )
        print(f"Workbook: {OUTPUT_DIR / 'Justlife_CRM.xlsx'}")
        print(f"Notes: {OUTPUT_DIR / 'NOTES.txt'}")
        print(f"Screenshots: {SHOTS_DIR}")


###############################################################################
# Entry point                                                                  #
###############################################################################


async def main() -> None:
    scraper = JustlifeScraper()
    await scraper.run()
    scraper.export_outputs()
    scraper.print_summary()


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except RuntimeError:
        loop = asyncio.get_event_loop()
        loop.run_until_complete(main())
