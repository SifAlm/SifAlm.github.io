"""Excel packaging utilities for the Tamam Justlife deliverables."""
from __future__ import annotations

from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ZEBRA_FILL = PatternFill("solid", fgColor="F2F2F2")
QA_FILL = PatternFill("solid", fgColor="FFC7CE")
OFFER_FILL = PatternFill("solid", fgColor="C6EFCE")
GRADIENT_COLORS = ["FFF2CC", "FFEB84", "FFD966", "FFC000", "FF8C00"]


class ExcelBuilder:
    """Create the Justlife_CRM.xlsx workbook from the CSV dataframes."""

    def __init__(self, output_path: Path) -> None:
        self.output_path = output_path
        self.workbook = Workbook()
        # Remove the default sheet to control ordering explicitly.
        default_sheet = self.workbook.active
        self.workbook.remove(default_sheet)

    # Sheet helpers --------------------------------------------------------
    @staticmethod
    def _style_header(row) -> None:
        for cell in row:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def _freeze_and_filter(sheet: Worksheet) -> None:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    @staticmethod
    def _apply_zebra(sheet: Worksheet, start_row: int = 2) -> None:
        for idx, row in enumerate(sheet.iter_rows(min_row=start_row), start=start_row):
            if idx % 2 == 0:
                for cell in row:
                    if not cell.fill or cell.fill.fill_type is None:
                        cell.fill = ZEBRA_FILL

    def _write_dataframe(self, name: str, df: pd.DataFrame) -> Worksheet:
        sheet = self.workbook.create_sheet(title=name)
        rows = dataframe_to_rows(df, index=False, header=True)
        for row_index, row in enumerate(rows, start=1):
            sheet.append(row)
            if row_index == 1:
                self._style_header(sheet[row_index])
        self._freeze_and_filter(sheet)
        self._apply_zebra(sheet)
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = max(12, min(max_length + 2, 60))
        return sheet

    def _conditional_format_price_matrix(self, sheet: Worksheet, df: pd.DataFrame) -> None:
        if "grand_total" not in df.columns:
            return
        from openpyxl.formatting.rule import ColorScaleRule

        last_row = sheet.max_row
        last_col_letter = sheet.cell(row=1, column=sheet.max_column).column_letter
        rule = ColorScaleRule(start_type="min", start_color=GRADIENT_COLORS[0],
                              mid_type="percentile", mid_value=50, mid_color=GRADIENT_COLORS[2],
                              end_type="max", end_color=GRADIENT_COLORS[-1])
        sheet.conditional_formatting.add(f"A2:{last_col_letter}{last_row}", rule)

        if "qa_flags" in df.columns:
            for row_idx in range(2, last_row + 1):
                cell = sheet.cell(row=row_idx, column=df.columns.get_loc("qa_flags") + 1)
                if cell.value:
                    cell.fill = QA_FILL
        if "offer_state" in df.columns:
            for row_idx in range(2, last_row + 1):
                cell = sheet.cell(row=row_idx, column=df.columns.get_loc("offer_state") + 1)
                if str(cell.value).upper() == "ON":
                    for col in range(1, sheet.max_column + 1):
                        sheet.cell(row=row_idx, column=col).fill = OFFER_FILL

    def build_workbook(
        self,
        summary_df: pd.DataFrame,
        price_matrix_df: pd.DataFrame,
        addons_df: pd.DataFrame,
        availability_df: pd.DataFrame,
        offers_df: pd.DataFrame,
        pages_options_df: pd.DataFrame,
    ) -> Path:
        self._write_dataframe("Summary", summary_df)
        price_sheet = self._write_dataframe("Price Matrix", price_matrix_df)
        self._conditional_format_price_matrix(price_sheet, price_matrix_df)
        self._write_dataframe("Add-ons", addons_df)
        self._write_dataframe("Availability", availability_df)
        self._write_dataframe("Offers", offers_df)
        self._write_dataframe("Pages & Options", pages_options_df)
        self.workbook.save(self.output_path)
        return self.output_path

