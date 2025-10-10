import os
import sys
import json
import math
import time
import uuid
import queue
import types
import errno
import shutil
import random
import string
import typing as t
import datetime as dt
import dataclasses
from dataclasses import dataclass, field, asdict

# -- Logging helpers --------------------------------------------------------
def _log(level: str, message: str) -> None:
    timestamp = dt.datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')
    print(f"[{level}] {timestamp} | {message}")

# -- Environment & configuration -------------------------------------------
BASE_URL = os.getenv("BASE_URL", "https://www.justlife.com/en-AE")
DEFAULT_CITY = os.getenv("CITY", "Dubai")
try:
    MAX_COMBOS = int(os.getenv("MAX_COMBOS_PER_PAGE", "40"))
except ValueError:
    MAX_COMBOS = 40
try:
    DAYS_OF_SLOTS = int(os.getenv("DAYS_OF_SLOTS", "14"))
except ValueError:
    DAYS_OF_SLOTS = 14

OUTPUT_FILES = {
    "pages": "pages.csv",
    "options": "options.csv",
    "prices": "prices_raw.csv",
    "addons": "addons.csv",
    "availability": "availability.csv",
    "offers": "offers.csv",
    "elements": "elements.csv",
}
NOTES_PATH = "NOTES.txt"
EXCEL_PATH = "Justlife_CRM.xlsx"

# -- Installation of dependencies ------------------------------------------
REQUIRED_PACKAGES = [
    "requests",
    "beautifulsoup4",
    "lxml",
    "pandas",
    "openpyxl",
    "nest_asyncio",
]
OPTIONAL_PACKAGES = ["playwright"]


def ensure_packages(packages: t.List[str]) -> None:
    import subprocess
    for pkg in packages:
        try:
            __import__(pkg)
            _log("INFO", f"Package '{pkg}' already available")
        except Exception:
            _log("INFO", f"Installing package '{pkg}'")
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", pkg])
            except subprocess.CalledProcessError as exc:
                _log("WARN", f"Failed installing '{pkg}': {exc}")


ensure_packages(REQUIRED_PACKAGES)
ensure_packages(OPTIONAL_PACKAGES)

import requests
from bs4 import BeautifulSoup
import pandas as pd

try:
    import nest_asyncio
    nest_asyncio.apply()
except Exception as exc:
    _log("WARN", f"Failed applying nest_asyncio: {exc}")

try:
    from playwright.async_api import async_playwright
    PLAYWRIGHT_AVAILABLE = True
except Exception as exc:
    _log("WARN", f"Playwright not available: {exc}")
    PLAYWRIGHT_AVAILABLE = False

# -- Data classes -----------------------------------------------------------


def gen_uuid(namespace: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, namespace))


@dataclass
class PageRecord:
    page_id: str
    parent_page_id: str
    level: str
    url: str
    page_title_text: str
    detected_model_type: str
    city: str
    has_addons: str
    has_slots: str
    notes: str


@dataclass
class OptionRecord:
    page_id: str
    control_group: str
    option_code: str
    option_label: str
    option_type: str
    is_default: bool
    sort_index: int


@dataclass
class PriceRecord:
    page_id: str
    service_name: str
    model_type: str
    offer_state: str
    base_price: t.Optional[float]
    fees_total: t.Optional[float]
    vat_amount: t.Optional[float]
    vat_percent: t.Optional[float]
    grand_total: t.Optional[float]
    currency: str
    pricing_source_url: str
    collected_at: str
    qa_flags: str
    screenshot_ref: str
    dynamic_dims: t.Dict[str, t.Any] = field(default_factory=dict)


@dataclass
class AddonRecord:
    page_id: str
    addon_name: str
    addon_description: str
    addon_price: t.Optional[float]
    currency: str
    is_recommended: bool
    is_required: bool
    seen_on_step: str
    pricing_source_url: str
    collected_at: str


@dataclass
class AvailabilityRecord:
    page_id: str
    date_str: str
    time_slot_label: str
    slot_status: str
    lead_time_days: int
    collected_at: str


@dataclass
class OfferRecord:
    page_id: str
    offer_name: str
    offer_badge_text: str
    description: str
    eligibility: str
    discount_type: str
    discount_value: t.Optional[float]
    stacking_rules: str
    collected_at: str


@dataclass
class ElementRecord:
    page_id: str
    selector: str
    element_role: str
    raw_text: str
    normalized_text: str


# -- Helpers for HTTP -------------------------------------------------------


HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36",
    "Accept-Language": "en-US,en;q=0.9",
}


def fetch_url(url: str, retries: int = 3, timeout: int = 20) -> t.Optional[requests.Response]:
    for attempt in range(1, retries + 1):
        try:
            resp = requests.get(url, headers=HEADERS, timeout=timeout)
            if resp.status_code == 200:
                return resp
            else:
                _log("WARN", f"Request to {url} returned status {resp.status_code}")
        except requests.RequestException as exc:
            _log("WARN", f"Attempt {attempt} failed for {url}: {exc}")
            time.sleep(min(3, attempt))
    return None


# -- Discovery --------------------------------------------------------------


def parse_links_from_html(html: str, base_url: str) -> t.Set[str]:
    soup = BeautifulSoup(html, "lxml")
    links: t.Set[str] = set()
    for a in soup.select("a[href]"):
        href = a.get("href")
        if not href:
            continue
        if href.startswith("#"):
            continue
        if href.startswith("mailto:") or href.startswith("tel:"):
            continue
        url = requests.compat.urljoin(base_url, href)
        if "/blog" in url or "/legal" in url or "/privacy" in url:
            continue
        if not url.startswith(BASE_URL.rstrip("/")):
            continue
        links.add(url.split("#")[0])
    return links


def parse_sitemap(base_url: str) -> t.Set[str]:
    urls: t.Set[str] = set()
    sitemap_candidates = ["sitemap.xml", "sitemap_index.xml"]
    for suffix in sitemap_candidates:
        sitemap_url = requests.compat.urljoin(base_url, suffix)
        resp = fetch_url(sitemap_url)
        if not resp:
            continue
        soup = BeautifulSoup(resp.text, "xml")
        for loc in soup.select("loc"):
            link = loc.text.strip()
            if not link.startswith(BASE_URL.rstrip("/")):
                continue
            if "/blog" in link or "/legal" in link:
                continue
            urls.add(link.split("#")[0])
    return urls


# -- Model detection heuristics --------------------------------------------

MODEL_KEYWORDS = {
    "HOURS_PROS": ["hour", "professional", "cleaner"],
    "ROOM_MATERIAL": ["bedroom", "bathroom", "studio", "villa"],
    "APPLIANCE_COMBO": ["appliance", "fridge", "oven"],
    "UNIT_QUANTITY": ["unit", "piece", "sofa"],
    "AREA_BASED": ["sq", "sqm", "square"],
    "FLAT_PACKAGE": ["package", "plan", "basic", "standard", "premium"],
}


def detect_model(text: str) -> str:
    text_lower = text.lower()
    for model, keywords in MODEL_KEYWORDS.items():
        if any(keyword in text_lower for keyword in keywords):
            return model
    return "OTHER"


# -- Static parsing of page -------------------------------------------------


def extract_title(soup: BeautifulSoup) -> str:
    for selector in ["h1", "h2", "title"]:
        el = soup.select_one(selector)
        if el and el.get_text(strip=True):
            title = el.get_text(" ", strip=True)
            if title.startswith("http"):
                continue
            return title
    return ""


# Placeholder functions for options and pricing

def extract_options_static(soup: BeautifulSoup) -> t.List[OptionRecord]:
    options: t.List[OptionRecord] = []
    control_groups = soup.select("label")
    idx = 0
    for label in control_groups:
        text = label.get_text(" ", strip=True)
        if not text:
            continue
        option = OptionRecord(
            page_id="",
            control_group=text,
            option_code=text.lower().replace(" ", "_"),
            option_label=text,
            option_type="unknown",
            is_default=(idx == 0),
            sort_index=idx,
        )
        options.append(option)
        idx += 1
    return options


# -- Excel formatting -------------------------------------------------------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def autofit_columns(ws):
    for column_cells in ws.columns:
        length = 0
        column = column_cells[0].column
        for cell in column_cells:
            try:
                value = str(cell.value)
            except Exception:
                value = ""
            if value:
                length = max(length, len(value))
        ws.column_dimensions[get_column_letter(column)].width = min(max(length + 2, 10), 60)


def apply_table_style(ws):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill_color = "F2F2F2" if idx % 2 == 0 else "FFFFFF"
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=fill_color)


# -- Data collection orchestrator ------------------------------------------


def main():
    collected_at = dt.datetime.utcnow().replace(microsecond=0).isoformat() + "Z"

    pages: t.List[PageRecord] = []
    options: t.List[OptionRecord] = []
    prices: t.List[PriceRecord] = []
    addons: t.List[AddonRecord] = []
    availability: t.List[AvailabilityRecord] = []
    offers: t.List[OfferRecord] = []
    elements: t.List[ElementRecord] = []

    visited: t.Set[str] = set()
    to_visit: queue.Queue[str] = queue.Queue()

    base_resp = fetch_url(BASE_URL)
    if base_resp:
        soup = BeautifulSoup(base_resp.text, "lxml")
        base_links = parse_links_from_html(base_resp.text, BASE_URL)
        _log("INFO", f"Discovered {len(base_links)} links from base page")
        for link in base_links:
            to_visit.put(link)
    else:
        _log("ERROR", "Failed to load base URL; continuing with sitemap only")

    sitemap_links = parse_sitemap(BASE_URL)
    _log("INFO", f"Discovered {len(sitemap_links)} links from sitemap")
    for link in sitemap_links:
        to_visit.put(link)

    max_pages = 50
    processed_count = 0

    while not to_visit.empty() and processed_count < max_pages:
        url = to_visit.get()
        if url in visited:
            continue
        visited.add(url)
        processed_count += 1

        resp = fetch_url(url)
        if not resp:
            _log("WARN", f"Skipping {url} due to fetch failure")
            continue

        soup = BeautifulSoup(resp.text, "lxml")
        title = extract_title(soup) or ""
        detected_model = detect_model(soup.get_text(" ", strip=True))

        page_id = gen_uuid(url)
        page_record = PageRecord(
            page_id=page_id,
            parent_page_id="",
            level="service",
            url=url,
            page_title_text=title,
            detected_model_type=detected_model,
            city=DEFAULT_CITY,
            has_addons="unknown",
            has_slots="unknown",
            notes="Static fallback; dynamic data unavailable"
        )
        pages.append(page_record)

        elements.append(ElementRecord(
            page_id=page_id,
            selector="title",
            element_role="title",
            raw_text=title,
            normalized_text=title.lower(),
        ))

        page_options = extract_options_static(soup)
        for idx, opt in enumerate(page_options):
            opt.page_id = page_id
            opt.sort_index = idx
            options.append(opt)

        if not title or title.startswith("http"):
            notes = f"Title fallback used for {url}"
        else:
            notes = ""

        price_record = PriceRecord(
            page_id=page_id,
            service_name=title if title and not title.startswith("http") else "",
            model_type=detected_model,
            offer_state="UNKNOWN",
            base_price=None,
            fees_total=None,
            vat_amount=None,
            vat_percent=None,
            grand_total=None,
            currency="",
            pricing_source_url=url,
            collected_at=collected_at,
            qa_flags="NO_TOTAL",
            screenshot_ref="",
        )
        prices.append(price_record)

    # -- Write CSVs ---------------------------------------------------------
    os.makedirs(".", exist_ok=True)

    def write_csv(filename: str, df: pd.DataFrame) -> None:
        df.to_csv(filename, index=False)
        _log("INFO", f"Wrote {filename} with {len(df)} rows")

    pages_df = pd.DataFrame([asdict(p) for p in pages])
    if pages_df.empty:
        pages_df = pd.DataFrame(columns=[field.name for field in dataclasses.fields(PageRecord)])
    write_csv(OUTPUT_FILES["pages"], pages_df)

    options_df = pd.DataFrame([asdict(o) for o in options])
    if options_df.empty:
        options_df = pd.DataFrame(columns=[field.name for field in dataclasses.fields(OptionRecord)])
    write_csv(OUTPUT_FILES["options"], options_df)

    price_rows: t.List[dict] = []
    dimension_keys: t.Set[str] = set()
    for record in prices:
        row = asdict(record)
        dims = row.pop("dynamic_dims", {})
        dimension_keys.update(dims.keys())
        row.update(dims)
        price_rows.append(row)

    base_columns = [
        "page_id",
        "service_name",
        "model_type",
        "offer_state",
        "base_price",
        "fees_total",
        "vat_amount",
        "vat_percent",
        "grand_total",
        "currency",
        "pricing_source_url",
        "collected_at",
        "qa_flags",
        "screenshot_ref",
    ]
    price_columns = base_columns + sorted(dimension_keys)
    prices_df = pd.DataFrame(price_rows, columns=price_columns)
    write_csv(OUTPUT_FILES["prices"], prices_df)

    addons_df = pd.DataFrame([asdict(a) for a in addons])
    if addons_df.empty:
        addons_df = pd.DataFrame(columns=[field.name for field in dataclasses.fields(AddonRecord)])
    write_csv(OUTPUT_FILES["addons"], addons_df)

    availability_df = pd.DataFrame([asdict(a) for a in availability])
    if availability_df.empty:
        availability_df = pd.DataFrame(columns=[field.name for field in dataclasses.fields(AvailabilityRecord)])
    write_csv(OUTPUT_FILES["availability"], availability_df)

    offers_df = pd.DataFrame([asdict(o) for o in offers])
    if offers_df.empty:
        offers_df = pd.DataFrame(columns=[field.name for field in dataclasses.fields(OfferRecord)])
    write_csv(OUTPUT_FILES["offers"], offers_df)

    elements_df = pd.DataFrame([asdict(e) for e in elements])
    if elements_df.empty:
        elements_df = pd.DataFrame(columns=[field.name for field in dataclasses.fields(ElementRecord)])
    write_csv(OUTPUT_FILES["elements"], elements_df)

    # -- Create Excel workbook ---------------------------------------------
    wb = Workbook()

    def add_sheet_from_df(name: str, df: pd.DataFrame) -> None:
        ws = wb.create_sheet(title=name)
        if df.empty:
            ws.append(["No data"])
            return
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append(list(row))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        apply_table_style(ws)
        autofit_columns(ws)

    wb.remove(wb.active)

    summary_ws = wb.create_sheet(title="Summary")
    summary_ws.append(["Metric", "Value"])
    summary_data = {
        "#pages": len(pages_df),
        "#services": pages_df["page_title_text"].nunique() if not pages_df.empty else 0,
        "#price_combos": len(prices_df),
        "#addons": len(addons_df),
        "#availability_records": len(availability_df),
        "Earliest lead time": availability_df["lead_time_days"].min() if not availability_df.empty else "",
        "Average total": prices_df["grand_total"].mean() if not prices_df.empty else "",
        "City": DEFAULT_CITY,
    }
    for metric, value in summary_data.items():
        summary_ws.append([metric, value])
    summary_ws.freeze_panes = "A2"
    apply_table_style(summary_ws)
    autofit_columns(summary_ws)

    add_sheet_from_df("Pages", pages_df)
    add_sheet_from_df("Options", options_df)
    add_sheet_from_df("Pricing", prices_df)
    add_sheet_from_df("Addons", addons_df)
    add_sheet_from_df("Availability", availability_df)
    add_sheet_from_df("Offers", offers_df)
    add_sheet_from_df("Diagnostics", elements_df)

    wb.save(EXCEL_PATH)
    _log("INFO", f"Excel workbook saved to {EXCEL_PATH}")

    # -- Notes --------------------------------------------------------------
    with open(NOTES_PATH, "w", encoding="utf-8") as fh:
        fh.write("Static fallback run; dynamic browser not executed.\n")
        fh.write(f"Processed {len(pages)} pages.\n")
        fh.write("No totals captured; pricing requires dynamic rendering.\n")
    _log("INFO", f"Notes written to {NOTES_PATH}")

    # -- Summary ------------------------------------------------------------
    _log("INFO", "Run summary:")
    _log("INFO", json.dumps({
        "pages": len(pages_df),
        "options": len(options_df),
        "price_states": len(prices_df),
        "addons": len(addons_df),
        "availability": len(availability_df),
        "offers": len(offers_df),
        "outputs": list(OUTPUT_FILES.values()) + [EXCEL_PATH, NOTES_PATH],
    }, indent=2))


if __name__ == "__main__":
    try:
        import dataclasses
        main()
    except Exception as exc:
        _log("ERROR", f"Unhandled exception: {exc}")
        with open(NOTES_PATH, "a", encoding="utf-8") as fh:
            fh.write(f"Error occurred: {exc}\n")
        sys.exit(1)
